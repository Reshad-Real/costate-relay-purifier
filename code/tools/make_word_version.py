"""Word (.docx) version of the anonymised review manuscript (ISA Transactions prefers MS Word).

The LaTeX source manuscript2/main.tex stays the master. This script flattens it into plain LaTeX that pandoc can
convert: numbers from the generated macros, citation numbers and cross-reference numbers from the compiled main.aux,
the bibliography from main.bbl, equations numbered as in the PDF, theorem-like environments as bold run-in headings,
and every figure as a 300-dpi PNG of the exported vector figure. pandoc converts the result to .docx with native Word
equations; python-docx then sets A4, 2.5-cm margins, 12-pt Times New Roman, single spacing and line numbers.
Word itself (COM automation) renders the .docx to PDF to check the page count (limit 30) and the layout.

Run after build.sh (needs main.aux, main.bbl and the exported figures):
    python make_word_version.py
Output: <submission folder>/Manuscript_anonymised.docx (+ _word/Manuscript_anonymised_docx_render.pdf for checking)."""
import os, re, sys, subprocess, shutil
import fitz
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
MS = os.path.join(ROOT, "manuscript2")
GEN = os.path.join(ROOT, "manuscript", "generated")
SUB = os.path.join(os.path.dirname(ROOT), "1_LATEST_MANUSCRIPT_ISA_Transactions")
FIGDIR = os.path.join(SUB, "Figures")
WORK = os.path.join(MS, "_word")
PANDOC = os.environ.get("PANDOC", r"C:\Users\Asus ROG Strix G16\AppData\Local\Pandoc\pandoc.exe")
TEXTWIDTH_CM = 16.0          # A4 with 2.5-cm margins
SINGLE = 0.62                # \singlefig in the review build


# ------------------------------------------------------------------ helpers
def braced(s, i):
    """s[i] == '{': return (content, index after the matching '}')."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces at " + s[i:i + 60])


def strip_cmd_arg(s, cmd):
    """Remove \\cmd{...} (with balanced braces) everywhere and return (s, list of removed arguments)."""
    out, args, k = [], [], 0
    pat = re.compile(r"\\" + cmd + r"\s*\{")
    while True:
        m = pat.search(s, k)
        if not m:
            out.append(s[k:])
            break
        out.append(s[k:m.start()])
        arg, e = braced(s, m.end() - 1)
        args.append(arg)
        k = e
    return "".join(out), args


def envs(s, name):
    """Yield (start, end, body, optarg) for every \\begin{name}...\\end{name} (not nested)."""
    b, e = "\\begin{" + name + "}", "\\end{" + name + "}"
    k = 0
    while True:
        i = s.find(b, k)
        if i < 0:
            return
        j = s.find(e, i)
        body = s[i + len(b):j]
        opt = None
        m = re.match(r"\s*\[([^\]]*)\]", body)
        if m:
            opt, body = m.group(1), body[m.end():]
        yield i, j + len(e), body, opt
        k = j + len(e)


def replace_envs(s, name, fn):
    parts, k = [], 0
    for i, j, body, opt in list(envs(s, name)):
        parts.append(s[k:i]); parts.append(fn(body, opt)); k = j
    parts.append(s[k:])
    return "".join(parts)


def compress(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f"{nums[i]}\u2013{nums[j]}")
        else:
            out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ", ".join(out)


# ------------------------------------------------------------------ inputs
def load_macros():
    mac = {}
    for f in ("fallback.tex", "macros.tex"):            # macros.tex last: real values win
        p = os.path.join(GEN, f)
        if not os.path.exists(p):
            continue
        for line in open(p, encoding="utf-8"):
            m = re.match(r"\s*\\(?:newcommand|providecommand|renewcommand)\*?\{\\([A-Za-z]+)\}\{(.*)\}\s*$", line)
            if m:
                mac[m.group(1)] = m.group(2)
    t = open(os.path.join(MS, "title.tex"), encoding="utf-8").read()
    m = re.search(r"\\newcommand\{\\PaperTitle\}\{", t)
    mac["PaperTitle"] = braced(t, m.end() - 1)[0]
    return mac


def load_aux():
    a = open(os.path.join(MS, "main.aux"), encoding="utf-8", errors="ignore").read()
    labels = {k: v for k, v in re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", a)}
    cites = {k: int(v) for k, v in re.findall(r"\\bibcite\{([^}]+)\}\{\{?(\d+)", a)}    # plain or natbib form
    return labels, cites


def load_bbl():
    s = open(os.path.join(MS, "main.bbl"), encoding="utf-8", errors="ignore").read()
    items = re.split(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", s)[1:]
    refs = []
    for key, txt in zip(items[0::2], items[1::2]):
        txt = txt.split("\\end{thebibliography}")[0]
        txt = re.sub(r"\\newblock\s*", " ", txt)
        txt = re.sub(r"\s+", " ", txt).strip()
        refs.append((key, txt))
    return refs


# ------------------------------------------------------------------ flattening
def flatten():
    mac = load_macros()
    labels, cites = load_aux()
    src = open(os.path.join(MS, "main.tex"), encoding="utf-8").read()
    s = src[src.index("\\begin{document}") + len("\\begin{document}"):src.index("\\end{document}")]
    s = re.sub(r"(?<!\\)%.*", "", s)                                    # comments

    # review build: drop the \ifdefined\FINALVERSION branches, keep the \else branch
    while "\\ifdefined\\FINALVERSION" in s:
        i = s.index("\\ifdefined\\FINALVERSION")
        j = s.index("\\fi", i)
        blk = s[i + len("\\ifdefined\\FINALVERSION"):j]
        keep = blk.split("\\else", 1)[1] if "\\else" in blk else ""
        s = s[:i] + keep + s[j + 3:]

    # front matter: title, abstract, keywords
    fm_i, fm_j = s.index("\\begin{frontmatter}"), s.index("\\end{frontmatter}") + len("\\end{frontmatter}")
    fm = s[fm_i:fm_j]
    abstract = fm[fm.index("\\begin{abstract}") + 16:fm.index("\\end{abstract}")].strip()
    kw = fm[fm.index("\\begin{keyword}") + 15:fm.index("\\end{keyword}")].strip().replace("\\sep", ";")
    kw = re.sub(r"\s*;\s*", "; ", re.sub(r"\s+", " ", kw))
    s = s[:fm_i] + ("\\title{\\PaperTitle}\n\n\\section*{Abstract}\n" + abstract +
                    "\n\n\\textbf{Keywords:} " + kw + "\n\n") + s[fm_j:]

    # numbers and preamble shorthands
    def mrep(m):
        name = m.group(1)
        if name in mac:
            return mac[name]
        return m.group(0)
    for _ in range(3):                                                   # macros may contain macros
        s = re.sub(r"\\([A-Za-z]+)(?![A-Za-z])(?:\{\})?", mrep, s)
    s = s.replace("\\ensuremath{-}", "\u2212").replace("\\ensuremath", "")
    s = re.sub(r"\\U(?![A-Za-z])", r"\\mathcal{U}", s)
    s = re.sub(r"\\Z(?![A-Za-z])", r"\\mathcal{Z}", s)
    s = re.sub(r"\\P(?![A-Za-z])", r"\\mathcal{P}", s)

    # citations and cross-references
    def cite(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        return "[" + compress([cites[k] for k in keys]) + "]"
    s = re.sub(r"\\cite\{([^}]*)\}", cite, s)
    s = re.sub(r"\\eqref\{([^}]*)\}", lambda m: "(" + labels[m.group(1)] + ")", s)
    s = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels[m.group(1)], s)

    # equations: all numbered in order of appearance, like the PDF
    eqn = [0]
    def eq(body, opt):
        eqn[0] += 1
        body, labs = strip_cmd_arg(body, "label")
        for l in labs:                                                   # consistency with the PDF numbering
            assert labels[l] == str(eqn[0]), (l, labels[l], eqn[0])
        return "\n$$" + body.strip() + "\\qquad(" + str(eqn[0]) + ")$$\n"
    s = replace_envs(s, "equation", eq)

    # theorem-like environments as bold run-in headings with the PDF's numbers
    for env, word in [("theorem", "Theorem"), ("proposition", "Proposition"), ("lemma", "Lemma"),
                      ("remark", "Remark")]:
        def thm(body, opt, word=word):
            body, labs = strip_cmd_arg(body, "label")
            num = labels[labs[0]] if labs else "?"
            head = f"\\textbf{{{word} {num}}}" + (f" ({opt})" if opt else "") + ": "
            return "\n\n" + head + body.strip() + "\n\n"
        s = replace_envs(s, env, thm)
    s = replace_envs(s, "proof", lambda body, opt: "\n\n\\emph{Proof:} " + body.strip() + "\n\n")

    # figures: exported vector figures as PNG, numbered captions
    def figure(body, opt, wide):
        _, labs = strip_cmd_arg(body, "label")
        num = int(labels[labs[0]])
        m = re.search(r"\\caption\{", body)
        cap = braced(body, m.end() - 1)[0]
        w = TEXTWIDTH_CM if wide else TEXTWIDTH_CM * SINGLE
        if labs[0] == "fig:testbed":
            w = TEXTWIDTH_CM
        png = f"Figure_{num}.png"
        return (f"\n\n\\begin{{wordfig}}\n\\centering\n\\includegraphics[width={w:.1f}cm]{{{png}}}\n"
                f"\\caption{{Figure {num}: {cap}}}\n\\end{{wordfig}}\n\n")
    s = replace_envs(s, "figure*", lambda b, o: figure(b, o, True))
    s = replace_envs(s, "figure", lambda b, o: figure(b, o, False))

    # tables: generated tabulars inlined, numbered captions above
    def table(body, opt):
        _, labs = strip_cmd_arg(body, "label")
        num = labels[labs[0]]
        m = re.search(r"\\caption\{", body)
        cap = braced(body, m.end() - 1)[0]
        g = re.search(r"\\gentab\{([^}]*)\}", body)
        if g:
            tab = open(os.path.join(GEN, f"tab_{g.group(1)}.tex"), encoding="utf-8").read()
            tab = re.sub(r"\\ensuremath\{-\}", "\u2212", tab)
        else:
            i = body.index("\\begin{tabular}")
            tab = body[i:body.index("\\end{tabular}") + len("\\end{tabular}")]
        tab = re.sub(r"\\begin\{tabular\}\{[^\n]*?\}\n", lambda m: re.sub(r"L\{", "p{", m.group(0)), tab, count=1)
        tab = re.sub(r"\\\\(\s*)\[", r"\\\\\1{}[", tab)          # a cell starting with "[n]" is not a \\[skip]
        tab = re.sub(r"(\\(?:toprule|midrule|bottomrule|hline))(\s*)\[", r"\1\2{}[", tab)   # nor a rule width
        return (f"\n\n\\begin{{wordtab}}\n\\caption{{Table {num}: {cap}}}\n{tab}\n\\end{{wordtab}}\n\n")
    s = replace_envs(s, "table*", table)
    s = replace_envs(s, "table", table)

    # algorithm as a numbered list
    def algorithm(body, opt):
        _, labs = strip_cmd_arg(body, "label")
        num = labels[labs[0]] if labs else "1"
        m = re.search(r"\\caption\{", body)
        cap = braced(body, m.end() - 1)[0]
        steps = re.findall(r"\\State\s*(.*)", body)
        steps = [st.replace("\\gets", "\\leftarrow") for st in steps]
        return ("\n\n\\textbf{Algorithm " + num + ":} " + cap + "\n\n\\begin{enumerate}\n" +
                "".join("\\item " + st + "\n" for st in steps) + "\\end{enumerate}\n\n")
    s = replace_envs(s, "algorithm", algorithm)

    for a, b in [("wordfig", "figure"), ("wordtab", "table")]:
        s = s.replace("{" + a + "}", "{" + b + "}")

    # layout-only commands
    for c in ["centering", "footnotesize", "small", "FloatBarrier", "noindent", "hfill", "raggedbottom"]:
        s = re.sub(r"\\" + c + r"(?![A-Za-z])", "", s)

    # references from the compiled bibliography (same numbers as the PDF)
    refs = load_bbl()
    order = sorted(refs, key=lambda kv: cites[kv[0]])
    bib = "\n\n\\section*{References}\n\n" + "\n\n".join(f"[{cites[k]}] {t}" for k, t in order) + "\n"
    s = re.sub(r"\\bibliographystyle\{[^}]*\}", "", s)
    s = re.sub(r"\\bibliography\{[^}]*\}", lambda m: bib, s)
    return "\\documentclass{article}\n\\begin{document}\n" + s + "\n\\end{document}\n"


# ------------------------------------------------------------------ figures
def figures():
    for f in sorted(os.listdir(FIGDIR)):
        if f.startswith("Figure_") and f.endswith(".pdf"):
            d = fitz.open(os.path.join(FIGDIR, f))
            d[0].get_pixmap(dpi=300).save(os.path.join(WORK, f[:-4] + ".png"))


# ------------------------------------------------------------------ docx styling
def set_font(style, size=12, bold=None, italic=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def style_reference(path):
    d = Document(path)
    for st in d.styles:
        if st.type != 1:                                      # paragraph styles only
            continue
        n = st.name
        if n.startswith("Heading"):
            set_font(st, 12, bold=True, italic=False)
            st.paragraph_format.space_before = Pt(9); st.paragraph_format.space_after = Pt(3)
        elif n == "Title":
            set_font(st, 14, bold=True)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif n in ("Image Caption", "Table Caption", "Caption"):
            set_font(st, 11, italic=False)
            st.paragraph_format.space_before = Pt(3); st.paragraph_format.space_after = Pt(9)
        else:
            set_font(st, 12)
        pf = st.paragraph_format
        pf.line_spacing = 1.0
        if n in ("Normal", "Body Text", "First Paragraph", "Compact"):
            pf.space_before = Pt(0); pf.space_after = Pt(3)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if n != "Compact" else None
    d.save(path)


def finish(path):
    d = Document(path)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
        sec._sectPr.append(ln)
    # references: left-aligned (justified URLs stretch the lines)
    in_refs = False
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "References":
            in_refs = True
            continue
        if in_refs:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
    # tables: full width, widths fitted to the content, left-aligned 10-pt (9-pt for wide tables) text,
    # booktabs-like rules (top, below the header, bottom)
    for t in d.tables:
        ncol = len(t.columns)
        size = Pt(9) if ncol >= 8 else Pt(10)
        for row in t.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for w in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(w)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = size
        tblPr = t._tbl.tblPr
        for tag in ("w:tblW", "w:tblLayout"):
            for e in tblPr.findall(qn(tag)):
                tblPr.remove(e)
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), "5000"); tw.set(qn("w:type"), "pct"); tblPr.append(tw)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "autofit"); tblPr.append(lay)
        for e in tblPr.findall(qn("w:tblCellMar")):
            tblPr.remove(e)
        mar = OxmlElement("w:tblCellMar")                          # narrow cell padding (Word default 0.19 cm)
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), "60"); m.set(qn("w:type"), "dxa"); mar.append(m)
        tblPr.append(mar)
        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for gc in grid.findall(qn("w:gridCol")):
                if gc.get(qn("w:w")) is not None:
                    del gc.attrib[qn("w:w")]
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
            b.append(e)
        tblPr.append(b)
        if t.rows:
            for cell in t.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
                tb.append(e); tcPr.append(tb)
    d.save(path)


def render_with_word(docx, pdf):
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx}', $false, $true); "
          f"$n = $d.ComputeStatistics(2); $d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit(); "
          f"Write-Output $n")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=600)
    return r.stdout.strip(), r.stderr.strip()


def main():
    os.makedirs(WORK, exist_ok=True)
    flat = flatten()
    open(os.path.join(WORK, "flat.tex"), "w", encoding="utf-8").write(flat)
    figures()
    ref = os.path.join(WORK, "reference.docx")
    subprocess.run([PANDOC, "-o", ref, "--print-default-data-file", "reference.docx"], check=True)
    style_reference(ref)
    out = os.path.join(WORK, "Manuscript_anonymised.docx")
    r = subprocess.run([PANDOC, "flat.tex", "-f", "latex", "-t", "docx", "--number-sections",
                        "--reference-doc", "reference.docx", "-o", out], cwd=WORK, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stderr); sys.exit(1)
    if r.stderr.strip():
        print("pandoc warnings:\n" + r.stderr.strip()[:3000])
    finish(out)
    pdf = os.path.join(WORK, "Manuscript_anonymised_docx_render.pdf")
    n, err = render_with_word(out, pdf)
    print(f"Word pages: {n}" + (f" (stderr: {err[:300]})" if err else ""))
    dst = os.path.join(SUB, "Manuscript_anonymised.docx")
    try:
        shutil.copy2(out, dst)
        print("written:", dst)
    except PermissionError:
        print("WARNING: could not overwrite", dst, "(open in Word?)")


if __name__ == "__main__":
    main()
